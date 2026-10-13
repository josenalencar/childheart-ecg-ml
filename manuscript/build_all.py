"""Build the Pediatric Cardiology submission package.

    python3 _source/build_all.py "<package folder>" <copy of ~/Zotero/zotero.sqlite>

Derived from the sibling JAMA Pediatrics builder (CHILDHEART LV score/JAMA Pediatrics
submission/_source/build_all.py); the docx helpers are reused almost verbatim. What differs is
the journal, and every difference is driven by _source/JOURNAL_RULES.md:

  * Times New Roman 12 pt, DOUBLE spaced, continuous line numbers and page numbers.
  * Citations are bracketed numerals in citation order -- [3], [1-3, 7] -- not superscripts,
    and the CSL style is Springer Basic (numeric, brackets), not JAMA.
  * No "Key Points" box (a JAMA device). Abstract is 150-250 words of prose and must carry a
    hypothesis; 4-6 keywords follow it.
  * "Statements and Declarations" replaces JAMA's "Article Information", with the seven items
    the journal returns submissions for omitting.
  * Figures are Fig1/Fig2, 174 mm wide, captions in the manuscript file as "Fig. 1".
  * The supplement is cited as "Online Resource 1" and shipped as ESM_1.pdf.

Zotero citations are inserted as live fields so the numbering and the bibliography are built by
Word, not by this script: the reference list renumbers itself if a citation is added. Note the
journal asks authors not to use field functions; the delivered copy is therefore produced by
unlinking the fields in Word (see README), while _source keeps the linked version.
"""
import re, sys, os, json, random, string, sqlite3, zipfile, shutil, html
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C, supp as X, checklists as K, zot
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
ZDB = sys.argv[2] if len(sys.argv) > 2 else '/private/tmp/zsrc_pedcard.sqlite'
ZUSER = 11286533
ZVER = "9.0.6"
ZSTYLE = "http://www.zotero.org/styles/springer-basic-brackets"
random.seed(20261013)

# ------------------------------------------------------------------ citation order
def scan(t):
    """Citation keys in first-appearance order; [[a, b]] is two keys, not one."""
    out = []
    for grp in re.findall(r'\[\[([^\]]+)\]\]', t):
        out.extend(k.strip() for k in grp.split(',') if k.strip())
    return out

order = []
def note(keys):
    for k in keys:
        k = k.strip()
        if k not in order:
            if k not in zot.KEYS: raise KeyError(f'build_all: unknown citation key {k!r}')
            order.append(k)

note(scan(C.ABSTRACT))
for _, ps in C.MAIN:
    for p in ps: note(scan(p))
for _, t in C.BACK: note(scan(t))
for t in X.ALL_TEXT: note(scan(t))
NUM = {k: i + 1 for i, k in enumerate(order)}

def compress(nums, dash="-"):
    nums = sorted(nums); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(str(nums[i]) if j == i else
                   (f'{nums[i]}, {nums[j]}' if j == i + 1 else f'{nums[i]}{dash}{nums[j]}'))
        i = j + 1
    return ", ".join(out)

# ------------------------------------------------------------------ Zotero CSL
CSL = {}
db = sqlite3.connect(ZDB)
for k in order:
    CSL[k] = zot.csl(db, zot.KEYS[k])[0::2]

# Journal titles that are ALL CAPS or abbreviated oddly in the library. Corrected in this
# embedded copy only -- the library items themselves are left alone (they are the user's
# curation), so nothing here edits Zotero.
FIX = {
    'peguero':      {'container-title': 'Journal of the American College of Cardiology'},
    'nussinovitch': {'container-title': 'Journal of Clinical Rheumatology'},
    'devries':      {'container-title': 'Acta Obstetricia et Gynecologica Scandinavica'},
    'pachiyannan':  {'container-title': 'Technologies'},
}
for k, patch in FIX.items():
    if k in CSL: CSL[k][1].update(patch)

def rid(n=8):
    return "".join(random.choice(string.ascii_letters + string.digits) for _ in range(n))

# ------------------------------------------------------------------ docx helpers
def new_doc(line_numbers=False):
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    sec = d.sections[0]
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Inches(1))
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        e = OxmlElement('w:fldChar' if kind else 'w:instrText')
        if kind: e.set(qn('w:fldCharType'), kind)
        else: e.set(qn('xml:space'), 'preserve'); e.text = txt
        r._r.append(e)
    if line_numbers:   # the journal asks for continuous line numbering
        ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
        sec._sectPr.append(ln)
    return d

def fld(p, kind):
    r = p.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); r._r.append(e); return r

def instr(p, text):
    r = p.add_run(); e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = text; r._r.append(e); return r

def zotero_cite(p, keys, size=None):
    """A live Zotero field rendering [n] or [n-m]; Springer style, so no superscript."""
    nums = [NUM[k] for k in keys]
    keys = [k for _, k in sorted(zip(nums, keys))]
    shown = '[' + compress(nums, "-") + ']'
    payload = {"citationID": rid(),
               "properties": {"formattedCitation": shown, "plainCitation": shown, "noteIndex": 0},
               "citationItems": [{"id": CSL[k][0],
                                  "uris": [f"http://zotero.org/users/{ZUSER}/items/{zot.KEYS[k]}"],
                                  "itemData": CSL[k][1]} for k in keys],
               "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"}
    fld(p, 'begin'); instr(p, " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(payload, ensure_ascii=False) + " "); fld(p, 'separate')
    r = p.add_run(shown)
    if size: r.font.size = Pt(size)
    fld(p, 'end')

def add_runs(p, text, bold=False, italic=False, size=None, zotero=True):
    for part in re.split(r'(\[\[[^\]]+\]\]|\*\*[^*]+\*\*)', text):
        if not part: continue
        if part.startswith('[['):
            keys = [k.strip() for k in part[2:-2].split(',')]
            if zotero: zotero_cite(p, keys, size)
            else: p.add_run('[' + compress([NUM[k] for k in keys]) + ']')
            continue
        b = bold
        if part.startswith('**'): part, b = part[2:-2], True
        r = p.add_run(part); r.bold = b; r.italic = italic
        if size: r.font.size = Pt(size)

def para(d, text='', bold=False, italic=False, size=None, align=None, spacing=2.0, after=0, zotero=True):
    p = d.add_paragraph(); p.paragraph_format.line_spacing = spacing; p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    if text: add_runs(p, text, bold, italic, size, zotero)
    return p

def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def table(d, title, head, rows, note_text, size=9, widths=None):
    if title: para(d, title, bold=True, spacing=1.15, after=4)
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, hh in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(hh); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            for j, part in enumerate(re.split(r'\^\{([^}]*)\}', str(v))):
                r = cells[i].paragraphs[0].add_run(part); r.font.size = Pt(size); r.font.superscript = bool(j % 2)
                if i == 0 and not any(str(x).strip() for x in row[1:]): r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    if note_text:
        p = para(d, spacing=1.15, after=12)
        for j, part in enumerate(re.split(r'\^\{([^}]*)\}', note_text)):
            r = p.add_run(part); r.font.size = Pt(9); r.font.superscript = bool(j % 2)
    return t

def wc(t):
    return 0 if t.startswith('##') else len(re.sub(r'\[\[[^\]]+\]\]', '', t).replace('**', '').split())

def add_zotero_prefs(path):
    """ZOTERO_PREF_n custom document properties, so Word knows the style without being told."""
    data = (f'<data data-version="3" zotero-version="{ZVER}"><session id="{rid()}"/>'
            f'<style id="{ZSTYLE}" locale="en-US" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
            '<prefs><pref name="fieldType" value="Field"/>'
            '<pref name="automaticJournalAbbreviations" value="true"/></prefs></data>')
    chunks = [data[i:i + 255] for i in range(0, len(data), 255)]
    props = "".join(f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}" name="ZOTERO_PREF_{i + 1}">'
                    f'<vt:lpwstr>{html.escape(c)}</vt:lpwstr></property>' for i, c in enumerate(chunks))
    custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties '
              'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              + props + '</Properties>')
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            b = zin.read(it.filename)
            if it.filename == "[Content_Types].xml" and b"custom-properties" not in b:
                b = b.replace(b"</Types>", b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
            if it.filename == "_rels/.rels" and b"custom-properties" not in b:
                b = b.replace(b"</Relationships>", b'<Relationship Id="rIdZoteroCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
            if it.filename == "docProps/custom.xml": continue
            zout.writestr(it, b)
        zout.writestr("docProps/custom.xml", custom)
    shutil.move(tmp, path)

# ================================================================== 1. the manuscript
d = new_doc(line_numbers=True)

# -- title page (the journal wants acknowledgements and contributions here)
para(d, C.TITLE, bold=True, size=14, spacing=1.5, after=12)
para(d, C.AUTHORS, spacing=1.5, after=12)
for a in C.AFFILIATIONS: para(d, a, spacing=1.15)
para(d, '', spacing=1.15)
para(d, C.CORRESPONDENCE, spacing=1.15, after=12)

abs_words = wc(C.ABSTRACT)
main_words = sum(wc(p) for _, ps in C.MAIN for p in ps)
para(d, "Article type: Original Article (submitted as Research). Reporting guidelines: "
        "TRIPOD+AI and STARD 2015 (checklists uploaded).", spacing=1.15)
para(d, f"Word count: abstract, {abs_words} (limit 150-250); main text excluding headings, "
        f"{main_words}. Tables: {X.N_TABLES}. Figures: {X.N_FIGURES}. "
        f"References: {len(order)}. Supplementary material: Online Resource 1.", spacing=1.15)
para(d, "Acknowledgements: none.", spacing=1.15, after=12)
page_break(d)

# -- abstract and keywords
para(d, 'Abstract', bold=True)
para(d, C.ABSTRACT)
para(d, '', spacing=1.15)
p = para(d, spacing=1.5); r = p.add_run('Keywords  '); r.bold = True
p.add_run(' · '.join(C.KEYWORDS))
page_break(d)

# -- main text
for h, ps in C.MAIN:
    para(d, h, bold=True)
    for t in ps:
        para(d, t[2:], bold=True) if t.startswith('##') else para(d, t)
page_break(d)

# -- declarations
para(d, 'Statements and Declarations', bold=True)
for h, t in C.BACK:
    p = para(d, spacing=1.5, after=6); r = p.add_run(h + ': '); r.bold = True; add_runs(p, t)
page_break(d)

# -- references (one live ZOTERO_BIBL field spanning the list)
para(d, 'References', bold=True)
refps = []
for k in order:
    p = para(d, spacing=1.5, after=4, zotero=False)
    refps.append((p, f"{NUM[k]}. [{zot.KEYS[k]}] pending Zotero refresh"))
first, last = refps[0][0], refps[-1][0]
fld(first, 'begin')
instr(first, ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY ')
fld(first, 'separate')
for p, t in refps: p.add_run(t)
fld(last, 'end')
page_break(d)

# -- tables, then figure legends
for tb in X.TABLES:
    table(d, tb['title'], tb['head'], tb['rows'], tb['note'], widths=tb.get('widths'))
    page_break(d)
para(d, 'Figure Legends', bold=True)
for leg in X.FIGURE_LEGENDS:
    para(d, leg, spacing=2.0, after=12)

MS = os.path.join(OUT, "01 Manuscript.docx")
d.save(MS); add_zotero_prefs(MS)
print(f"wrote {MS}")
print(f"  abstract {abs_words} words ({'OK' if 150 <= abs_words <= 250 else 'OUT OF RANGE 150-250'}), "
      f"main text {main_words} words, {len(order)} references, "
      f"{X.N_TABLES} tables, {X.N_FIGURES} figures")
if not 150 <= abs_words <= 250:
    print("  !! fix the abstract length in content.py before submitting")

# ================================================================== 2. supplement (HTML -> PDF)
TRIPOD_CITE = ("Checklist item wording reproduced from Collins GS, Moons KGM, Dhiman P, et al. "
               "TRIPOD+AI statement: updated guidance for reporting clinical prediction models "
               "that use regression or machine learning methods. BMJ 2024;385:e078378 (CC BY 4.0).")
STARD_CITE = ("Checklist item wording reproduced from Bossuyt PM, Reitsma JB, Bruns DE, et al. "
              "STARD 2015: an updated list of essential items for reporting diagnostic accuracy "
              "studies. BMJ 2015;351:h5527 (CC BY-NC 4.0).")

def h2(t): return re.sub(r'\^\{([^}]*)\}', r'<sup>\1</sup>', html.escape(t))

SUPP_HEAD = "Online Resource 1. Supplementary Methods and Tables"
AUTH_LINE = ("Alencar Neto JN, Fernandes T, dos Santos Filho CA, et al. " + C.TITLE
             + ". Pediatr Cardiol. Submitted.")
css = ("body{font-family:'Times New Roman',serif;font-size:11pt;margin:28mm 20mm;}"
       "h1{font-size:14pt;} h2{font-size:12pt;margin-top:18pt;}"
       "table{border-collapse:collapse;width:100%;font-size:8.5pt;margin-bottom:6pt;}"
       "th,td{border:1px solid #444;padding:3px 5px;text-align:left;vertical-align:top;}"
       "th{background:#eee;} .note{font-size:8.5pt;margin:0 0 14pt 0;}"
       ".cite{font-style:italic;font-size:10pt;margin-bottom:16pt;}")
H = [f"<html><head><meta charset='utf-8'><style>{css}</style></head><body>",
     f"<h1>{h2(SUPP_HEAD)}</h1><div class='cite'>{h2(AUTH_LINE)}</div>"]
H.append("<h2>Contents</h2><ul>")
for t, _ in X.EMETHODS: H.append(f"<li>{h2(t)}</li>")
for e in X.ETABLES: H.append(f"<li>{h2(e['title'])}</li>")
H.append("</ul>")
for t, ps in X.EMETHODS:
    H.append(f"<h2>{h2(t)}</h2>")
    for q in ps: H.append(f"<p>{h2(q)}</p>")
for e in X.ETABLES:
    H.append(f"<h2>{h2(e['title'])}</h2><table><tr>"
             + "".join(f"<th>{h2(c)}</th>" for c in e['head']) + "</tr>")
    for row in e['rows']:
        H.append("<tr>" + "".join(f"<td>{h2(str(c))}</td>" for c in row) + "</tr>")
    H.append("</table>")
    if e.get('note'): H.append(f"<p class='note'>{h2(e['note'])}</p>")
H.append("</body></html>")
SUPP_HTML = os.path.join(os.path.dirname(os.path.abspath(__file__)), "supplement.html")
open(SUPP_HTML, "w", encoding="utf-8").write("\n".join(H))
print(f"wrote {SUPP_HTML} ({len(X.EMETHODS)} eMethods, {len(X.ETABLES)} eTables)")

# ================================================================== 3. checklists + cover letter
IT = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'items.json')))

def checklist(name, cite, rows, loc, fname):
    """One checklist docx.

    items.json stores each checklist as ragged rows straight from the published table: a
    one-element row is a section banner, a four-element row carries its own section label, and
    a three-element row continues the section above it. Normalising here means the item numbers
    line up with `loc`, which maps an item to where this manuscript reports it. An item with no
    entry prints blank rather than a guess, and the unlocated ones are listed on stdout so they
    cannot be forgotten.
    """
    k = new_doc()
    para(k, name, bold=True, size=14, spacing=1.3, after=4)
    para(k, cite, italic=True, spacing=1.15, after=6, zotero=False)
    para(k, AUTH_LINE, italic=True, spacing=1.15, after=12, zotero=False)
    body, missing, section = [], [], ''
    for r in rows[1:]:
        r = [str(x).strip() for x in r]
        if len(r) == 1:
            section = r[0]; body.append([section, '', '']); continue
        if len(r) == 4:
            section, item, text = (r[0] or section), r[1], r[3]
        elif len(r) == 3 and (r[0] == '' or not r[0][0].isdigit()):
            section, item, text = (r[0] or section), r[1], r[2]   # STARD: ['', no, item]
        else:
            item, text = r[0], r[-1]                              # TRIPOD continuation row
        if not item: continue
        where = loc.get(item, '')
        if not where: missing.append(item)
        body.append([item, text, where])
    table(k, '', ['Item', 'Checklist item', 'Reported on'], body, '', size=8,
          widths=[0.6, 4.2, 1.7])
    pth = os.path.join(OUT, fname); k.save(pth)
    n_items = len([x for x in body if x[0] and x[1]])
    print(f"wrote {pth} ({n_items} items, {n_items - len(missing)} located)")
    if missing:
        print(f"  NOT YET LOCATED: {', '.join(missing)}")

checklist("TRIPOD+AI Checklist", TRIPOD_CITE, IT['tripod'], K.TRIPOD_LOC,
          "05 TRIPOD+AI Checklist.docx")
checklist("STARD 2015 Checklist", STARD_CITE, IT['stard'], K.STARD_LOC,
          "06 STARD 2015 Checklist.docx")

cl = new_doc()
for t in K.COVER: para(cl, t, spacing=1.5, after=10, zotero=False)
CLP = os.path.join(OUT, "07 Cover Letter.docx")
cl.save(CLP); print(f"wrote {CLP}")

print("\ncitation order:", ", ".join(f"{NUM[k]}={k}" for k in order))
print("\nNEXT: open '01 Manuscript.docx' in Word and run Zotero > Refresh "
      "(style: Springer Basic, numeric, brackets), then unlink the fields for the copy you upload.")
